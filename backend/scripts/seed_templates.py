"""
One-time setup script: writes default templates to backend/data/templates/.

Re-running this OVERWRITES the named template(s) back to the standard
structure, discarding any manual edits made through the app - pass a specific
report type ("deals" or "pipeline") to limit the blast radius, e.g.:
    python scripts/seed_templates.py pipeline

Both "deals" and "pipeline" are single templates shared across all four
quarters, since the column layout and formulas don't change quarter to
quarter - only the close-date range being reported on does. Each title row
carries a {QUARTER} placeholder that process_quarter() fills in with the
selected quarter's label at processing time.

Deals formula convention:
  Difference of Sanction Amount and Disbursed Amount = Sanctioned - Disbursed
  Difference of Requested Loan Amount and Sanctioned Amount = Requested - Sanctioned
  % Requested vs Sanctioned = (Sanctioned - Requested) / Requested
  % Sanctioned vs Disbursed = (Disbursed - Sanctioned) / Sanctioned

Pipeline formula convention (deals still in progress - Sanction/Login/
Disbursement stages, not yet closed):
  Difference of Requested Amount Vs Login Amount = Requested - Login
  Percentage of Requested Amount Vs Login Amount = Login / Requested
  Both are only meaningful once Login Amount is known - the engine itself
  (excel_engine.build_output) leaves them blank for a row when Login Amount
  is missing, so the template doesn't need to do anything special for that.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from app.config import REPORT_TYPES, template_path  # noqa: E402

RUPEE_FORMAT = '"₹"#,##,##0.00'

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
NOTE_FONT = Font(italic=True, size=9, color="808080")

THIN = Side(style="thin", color="B7C0CC")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")
DATA_ALIGN = Alignment(vertical="center", horizontal="center", wrap_text=True)

NOTE_TEXT = (
    "How to add a new column: in row 2 add a header for the new column, and in row 3 "
    "(this example row) either enter a formula referencing other cells in row 3 (e.g. =E3-F3) "
    "for a calculated column, or leave it as an example value for a column that should be copied "
    "directly from the uploaded raw file (the header text must match the raw file's column exactly). "
    "Save and re-upload this template through the app."
)

DEALS_COLUMNS = [
    # (header, formula-or-None, number_format)
    ("Deal Name", None, "General"),
    ("Loan Type.", None, "General"),
    ("Financial Institute Name", None, "General"),
    ("Requested Loan Amount", None, RUPEE_FORMAT),
    ("Sanctioned Loan Amount ", None, RUPEE_FORMAT),
    ("Disbursed Loan Amount ", None, RUPEE_FORMAT),
    ("Balance Disbursement Amount ", "=E3-F3", RUPEE_FORMAT),
    ("Difference of Sanction Amount and Disbursed Amount", "=E3-F3", RUPEE_FORMAT),
    ("Difference of Requested Loan Amount and Sanctioned Amount ", "=D3-E3", RUPEE_FORMAT),
    (
        "Percentage of Requested Loan amount Vs Sanctinoned Loan Amount (In %)",
        "=(E3-D3)/D3",
        "0.00%",
    ),
    (
        "Percentage of Sanctioned Loan amount Vs Disbursed Loan Amount (in %)",
        "=(F3-E3)/E3",
        "0.00%",
    ),
    ("Date Of Disbursement ", None, "dd-mm-yyyy"),
    ("Lead Source", None, "General"),
    ("Status", None, "General"),
    ("Assigned To", None, "General"),
    ("Service Region", None, "General"),
    ("Service State", None, "General"),
]

DEALS_EXAMPLE_VALUES = {
    "Deal Name": "Example Deal",
    "Loan Type.": "Home Loan - New Purchase",
    "Financial Institute Name": "Example Bank Ltd",
    "Requested Loan Amount": 1000000,
    "Sanctioned Loan Amount ": 950000,
    "Disbursed Loan Amount ": 900000,
    "Date Of Disbursement ": "2026-04-15",
    "Lead Source": "Partner",
    "Status": "Closed Won",
    "Assigned To": "Example Person",
    "Service Region": "Chennai",
    "Service State": "Tamil Nadu",
}

PIPELINE_COLUMNS = [
    ("Deal Name", None, "General"),
    ("Stage", None, "General"),
    ("Probablity", None, "General"),
    ("Loan Type", None, "General"),
    ("Requested Loan Amount", None, RUPEE_FORMAT),
    ("Login Amount", None, RUPEE_FORMAT),
    ("Difference of Requested Amount Vs Login Amount", "=E3-F3", RUPEE_FORMAT),
    ("Percentage of Requested Amount Vs Login Amount", "=F3/E3", "0.00%"),
    ("Lead Source", None, "General"),
    ("Assigned To", None, "General"),
    ("Service Region", None, "General"),
]

PIPELINE_EXAMPLE_VALUES = {
    "Deal Name": "Example Deal",
    "Stage": "Login",
    "Probablity": "for Login it should be 50%",
    "Loan Type": "Home Loan - New Purchase",
    "Requested Loan Amount": 3000000,
    "Login Amount": 2800000,
    "Lead Source": "Partner",
    "Assigned To": "Example Person",
    "Service Region": "Chennai",
}


def build_template(sheet_title: str, title_text: str, columns: list, example_values: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    n_cols = len(columns)

    # title row - {QUARTER} is substituted with the selected quarter's label at process time
    ws.cell(1, 1, title_text)
    ws.cell(1, 1).font = TITLE_FONT
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    ws.row_dimensions[1].height = 26

    # header row
    for idx, (header, _formula, _fmt) in enumerate(columns, start=1):
        cell = ws.cell(2, idx, header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = CELL_BORDER
        ws.column_dimensions[get_column_letter(idx)].width = 24
    ws.row_dimensions[2].height = 48

    # example row (row 3)
    for idx, (header, formula, fmt) in enumerate(columns, start=1):
        cell = ws.cell(3, idx)
        cell.value = formula if formula is not None else example_values.get(header)
        cell.number_format = fmt
        cell.alignment = DATA_ALIGN
        cell.border = CELL_BORDER
    ws.row_dimensions[3].height = 20

    # usage note
    ws.cell(5, 1, NOTE_TEXT).font = NOTE_FONT
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=n_cols)
    ws.row_dimensions[5].height = 30
    ws.cell(5, 1).alignment = Alignment(wrap_text=True, vertical="top")

    return wb


def build_deals_template() -> Workbook:
    return build_template("DEALS", "DEALS For {QUARTER}", DEALS_COLUMNS, DEALS_EXAMPLE_VALUES)


def build_pipeline_template() -> Workbook:
    return build_template(
        "PIPELINE", "Pipeline Deals for {QUARTER}", PIPELINE_COLUMNS, PIPELINE_EXAMPLE_VALUES
    )


BUILDERS = {
    "deals": build_deals_template,
    "pipeline": build_pipeline_template,
}


def main():
    targets = sys.argv[1:] or REPORT_TYPES
    for report_type in targets:
        if report_type not in BUILDERS:
            print(f"skipping unknown report type: {report_type}")
            continue
        out_path = template_path(report_type)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        wb = BUILDERS[report_type]()
        wb.save(out_path)
        print(f"wrote {out_path}")


if __name__ == "__main__":
    main()
